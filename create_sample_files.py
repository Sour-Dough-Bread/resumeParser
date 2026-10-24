import os
import docx
from backend.sample_data import SAMPLE_RESUME_TEXT

os.makedirs("examples", exist_ok=True)

# 1. Save TXT sample
txt_path = os.path.join("examples", "sample_resume.txt")
with open(txt_path, "w", encoding="utf-8") as f:
    f.write(SAMPLE_RESUME_TEXT.strip())

# 2. Save DOCX sample using python-docx
docx_path = os.path.join("examples", "sample_resume.docx")
doc = docx.Document()
doc.add_heading("Alex Morgan", level=0)
doc.add_paragraph("San Francisco, CA | (555) 342-8901 | alex.morgan@email.com | linkedin.com/in/alexmorgan-dev")

for line in SAMPLE_RESUME_TEXT.strip().split("\n"):
    line = line.strip()
    if not line:
        continue
    if line in ["SUMMARY", "EDUCATION", "EXPERIENCE", "PROJECTS", "TECHNICAL SKILLS"]:
        doc.add_heading(line, level=1)
    elif line.startswith("- "):
        doc.add_paragraph(line[2:], style='List Bullet')
    elif line.startswith("Alex Morgan") or "linkedin.com" in line:
        continue
    else:
        doc.add_paragraph(line)

doc.save(docx_path)
print(f"Generated sample files:\n- {txt_path}\n- {docx_path}")
